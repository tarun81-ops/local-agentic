from pathlib import Path

from pptx import Presentation

from app.core.chunking import Chunk, chunk_located


def _shape_text(shape) -> list[str]:
    parts = []
    if shape.has_text_frame:
        parts.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            parts.append(" | ".join(c.text for c in row.cells if c.text.strip()))
    for child in getattr(shape, "shapes", []):  # grouped shapes
        parts.extend(_shape_text(child))
    return parts


def parse_pptx(path: Path) -> list[Chunk]:
    """Slide text, tables and speaker notes, cited by 1-indexed slide number."""
    chunks: list[Chunk] = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        chunks.extend(chunk_located("slide", i, "\n".join(p for p in parts if p.strip())))
    return chunks
