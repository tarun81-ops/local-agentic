"""A synthetic "personal documents" folder for the retrieval eval (run_eval.py --retrieval).

About 30 files in the kinds of folders people keep, with near-duplicates on purpose (three
invoices, two Everest meeting notes, a bank statement that also mentions the rent) so
ranking has to pick the right one. Written fresh into a temp folder on every run; nothing
here is real data.

    .venv\\Scripts\\python.exe eval\\retrieval_corpus.py <folder>     # just write the files"""
import sys
from email.message import EmailMessage
from pathlib import Path

# path -> content. str for text formats and PDFs (pages split on "\f"), list of rows for
# .xlsx, list of (title, body) for .pptx; .docx paragraphs are split on blank lines.
DOCS: dict[str, object] = {
    "Finance/Invoices/invoice_acme_2024-03.pdf": (
        "INVOICE No. 1042\nDate: 2 February 2024\nBill to: Acme Corp, 14 Harbour Road, Mumbai\n"
        "Description: Website redesign and content migration\nAmount due: $4,250\n"
        "Due date: 3 March 2024\nPayment terms: Net 30. Please pay by bank transfer."
    ),
    "Finance/Invoices/invoice_globex_2024-04.pdf": (
        "INVOICE No. 1043\nDate: 13 March 2024\nBill to: Globex Ltd, 2 Park Street, Kolkata\n"
        "Description: Logo and brand guidelines\nAmount due: $2,980\n"
        "Due date: 12 April 2024\nPayment terms: Net 30. Please pay by bank transfer."
    ),
    "Finance/Invoices/invoice_initech_2024-05.pdf": (
        "INVOICE No. 1044\nDate: 5 May 2024\nBill to: Initech Pvt Ltd, Hinjewadi Phase 2, Pune\n"
        "Description: Server migration to the cloud\nAmount due: $7,600\n"
        "Due date: 20 May 2024\nPayment terms: Net 15. Please pay by bank transfer."
    ),
    "Finance/electricity_bill_aug_2024.pdf": (
        "Tata Power - Electricity Bill\nConsumer No. 900012345\nBilling period: 1 Jul 2024 to 31 Jul 2024\n"
        "Units consumed: 412 kWh\nAmount payable: Rs. 3,540\nDue date: 15 August 2024\n"
        "Pay online before the due date to avoid a late fee of Rs. 50."
    ),
    "Finance/home_loan_sanction_letter.pdf": (
        "SANCTION LETTER\nDear Mr. Rohan Mehta,\nWe are pleased to sanction a home loan of Rs. 45,00,000 "
        "for the purchase of Plot 7, Lake View Residency, Pune.\nRate of interest: 8.60% p.a. floating, "
        "linked to the repo rate.\nTenure: 240 months\nEMI: Rs. 39,340\nProcessing fee: Rs. 10,000 plus GST.\n"
        "This sanction is valid for 6 months from the date of this letter."
    ),
    "Finance/Tax/itr_2023-24_summary.docx": (
        "Income Tax Return summary, Assessment Year 2024-25\n\nGross total income: Rs. 18,40,000\n\n"
        "Deductions under section 80C: Rs. 1,50,000 (PPF and ELSS)\n\nTax paid through TDS: Rs. 2,12,000\n\n"
        "Refund due: Rs. 18,200\n\nFiled on 28 July 2024. Acknowledgement number 123456789012345."
    ),
    "Finance/bank_statement_june_2024.xlsx": [
        ["Date", "Description", "Debit", "Credit", "Balance"],
        ["01-06-2024", "Salary credit Initech Pvt Ltd", "", "1,52,000", "2,10,450"],
        ["02-06-2024", "Rent transfer to R. Kulkarni", "32,000", "", "1,78,450"],
        ["05-06-2024", "SIP Axis Bluechip Fund", "10,000", "", "1,68,450"],
        ["09-06-2024", "Swiggy order", "640", "", "1,67,810"],
        ["14-06-2024", "Tata Power bill", "3,540", "", "1,64,270"],
    ],
    "Finance/budget_2025.xlsx": [
        ["Category", "Monthly amount (Rs.)"],
        ["Rent", "32,000"], ["Groceries", "12,000"], ["Utilities", "5,000"],
        ["SIP investments", "15,000"], ["Travel", "8,000"], ["Emergency fund", "10,000"],
    ],
    "Finance/mutual_fund_statement_q2.xlsx": [
        ["Fund", "Units", "NAV", "Value (Rs.)"],
        ["Axis Bluechip Fund - Growth", "1,204.55", "58.21", "70,117"],
        ["Parag Parikh Flexi Cap - Growth", "812.10", "76.40", "62,044"],
        ["HDFC Index Fund Nifty 50", "2,015.00", "212.35", "4,27,885"],
    ],
    "Work/Everest/everest_kickoff_meeting_notes.docx": (
        "Project Everest kickoff, 8 January 2024\n\nAttendees: Anita (PM), Rahul, Sofia\n\n"
        "Budget approved: $120,000 for phase one (discovery and prototype). Phase one ends 31 March.\n\n"
        "Risks: data access from the client's legacy CRM.\n\n"
        "Action items: Rahul to set up the staging environment; Sofia to draft the data-sharing agreement."
    ),
    "Work/Everest/everest_phase2_review.docx": (
        "Everest phase two review, 15 April 2024\n\n"
        "Phase two (integration and rollout) is delayed by six weeks because the CRM vendor changed.\n\n"
        "New vendor: Zenith Systems. Revised phase two budget: $85,000.\n\nGo-live moved to 30 August."
    ),
    "Work/apollo_status_update.pptx": [
        ("Project Apollo - Status, June 2024", "Owner: Rohan Mehta"),
        ("Timeline", "Design complete. Build 60% done. Launch planned for 1 October."),
        ("Risks", "Two engineers on leave in August. Payment gateway certification still pending."),
        ("Asks", "Approval for one contract QA engineer for eight weeks."),
    ],
    "Work/q3_sales_report.xlsx": [
        ["Region", "Target (USD)", "Actual (USD)"],
        ["North", "1,200,000", "1,350,000"], ["South", "900,000", "820,000"],
        ["West", "1,500,000", "1,610,000"], ["East", "700,000", "660,000"],
    ],
    "Work/onboarding_checklist.md": (
        "# New joiner checklist\n\n- Collect laptop from IT (Dell Latitude 7440)\n"
        "- Install GlobalProtect and connect to vpn.initech.example\n- Set up Okta MFA on your phone\n"
        "- Join Slack channels #general and #data-team\n- Complete POSH training within 30 days\n"
        "- Book a 1:1 with your manager in week one\n"
    ),
    "Work/team_offsite_agenda.txt": (
        "Data team offsite - Goa, 12-14 September 2024\n\n"
        "Day 1: arrive by 3 pm, welcome dinner at Fisherman's Wharf.\n"
        "Day 2: roadmap workshop, 2025 OKRs, beach volleyball in the evening.\n"
        "Day 3: retro, checkout by 11 am.\n\nThe travel desk books flights; send your preferences by 20 August.\n"
    ),
    "Work/performance_review_2024.docx": (
        "Annual performance review 2024 - Rohan Mehta\n\nRating: Exceeds expectations\n\n"
        "Strengths: led the Apollo reporting pipeline; mentored two interns.\n\n"
        "Development areas: stakeholder communication and delegating.\n\nRecommended increment: 12%."
    ),
    "Work/offer_letter_initech.pdf": (
        "Initech Pvt Ltd\nDear Rohan,\nWe are pleased to offer you the position of Senior Data Analyst.\n"
        "Annual CTC: Rs. 24,00,000\nJoining date: 1 July 2023\nNotice period: 60 days\n"
        "Probation: 6 months\nPlease sign and return a copy of this letter within 7 days."
    ),
    "Work/vendor_contract_zenith.docx": (
        "Master services agreement between Initech Pvt Ltd and Zenith Systems\n\n"
        "Term: 2 years from 1 May 2024.\n\nEither party may terminate with 90 days' written notice.\n\n"
        "Service level: 99.5% monthly uptime. Penalty: 5% of the monthly fee for each breach."
    ),
    "Personal/Priya Sharma CV.docx": (
        "Priya Sharma - Pune - priya.sharma@example.com\n\nExperience\n\n"
        "Data Analyst, Initech (2019-2023): built churn dashboards in Power BI; automated weekly "
        "reporting with Python.\n\nJunior Analyst, Globex (2017-2019)\n\n"
        "Education\n\nB.Tech Computer Science, COEP, 2017\n\nSkills: SQL, Python, Power BI, statistics"
    ),
    "Personal/Lease agreement B-12.pdf": (
        "LEAVE AND LICENSE AGREEMENT\nThis agreement is made on 25 July 2024 between Mr. R. Kulkarni "
        "(Licensor) and Mr. Rohan Mehta (Licensee) for Flat B-12, Green Park Society, Baner, Pune.\n"
        "License fee: Rs. 32,000 per month, payable by the 5th of each month.\n"
        "Security deposit: Rs. 1,50,000, refundable at the end of the term.\n"
        "Period: 11 months commencing 1 August 2024. Lock-in period: 6 months."
    ),
    "Personal/Insurance/car_insurance_policy_2024.pdf": (
        "MOTOR PACKAGE POLICY\nPolicy No. MOT/2024/778812\nVehicle: Hyundai Creta, MH12 AB 4521\n"
        "Insured Declared Value: Rs. 9,80,000\nPremium: Rs. 18,450 including GST\n"
        "Period of cover: 15 Feb 2024 to 14 Feb 2025\nAdd-ons: zero depreciation, roadside assistance."
    ),
    "Personal/Insurance/health_insurance_ecard.pdf": (
        "Star Health and Allied Insurance - e-card\nMembers: Rohan Mehta, Meera Mehta\n"
        "Plan: Family Floater\nSum insured: Rs. 10,00,000\nValid till: 31 March 2025\n"
        "TPA helpline: 1800-425-2255\nCashless treatment at network hospitals."
    ),
    "Personal/passport_renewal_checklist.txt": (
        "Passport re-issue - what to carry to the PSK appointment (Pune, 9 Nov, 10:30 am):\n"
        "- old passport, original and a self-attested copy of the first and last pages\n"
        "- address proof (Aadhaar or an electricity bill)\n- appointment receipt\n"
        "- fee already paid online: Rs. 1,500 (36 pages)\n"
    ),
    "Personal/Recipes/butter_chicken.md": (
        "# Butter chicken\n\nMarinate 500 g chicken in yoghurt, chilli powder and ginger-garlic paste "
        "for 2 hours.\nGrill or pan-sear, then simmer in a tomato, butter and cashew gravy for 20 minutes.\n"
        "Finish with kasuri methi and cream. Serves 4.\n"
    ),
    "Personal/emergency_contacts.csv": (
        "name,relation,phone\nMeera Mehta,spouse,+91 98200 11111\nDr. S. Rao,family doctor,+91 20 2567 8900\n"
        "Anil Mehta,father,+91 98220 22222\nSociety office,building,+91 20 2711 3344\n"
    ),
    "Personal/book_notes_atomic_habits.md": (
        "# Atomic Habits - notes\n\n- Get 1% better every day; small gains compound.\n"
        "- Habit stacking: after [current habit], I will [new habit].\n"
        "- Design the environment so the good choice is the easy one.\n- Never miss twice.\n"
    ),
    "Personal/gym_membership_receipt.pdf": (
        "Cult.fit - Payment receipt\nElite annual membership\nAmount paid: Rs. 22,999\n"
        "Valid from 1 Jan 2024 to 31 Dec 2024\nIncludes 2 pauses of up to 30 days each."
    ),
    "Personal/school_fee_receipt_term2.pdf": (
        "Vibgyor High - Fee receipt\nStudent: Aarav Mehta, Grade 3\nTerm 2 fees: Rs. 58,000\n"
        "Paid on 5 October 2024\nReceipt no. VH/24/5521"
    ),
    "Personal/fridge_warranty.pdf": (
        "Samsung 324 L frost-free refrigerator, model RT34\nPurchased 18 May 2023 from Croma\n"
        "1 year comprehensive warranty; 10 years on the digital inverter compressor.\n"
        "Service: 1800-40-7267864"
    ),
    "Personal/doctor_visit_notes.txt": (
        "Dr. S. Rao - 3 Sept 2024\nComplaint: seasonal allergy, sneezing, blocked nose.\n"
        "Advice: cetirizine 10 mg once at night for 5 days; saline nasal spray twice a day.\n"
        "Follow up if a fever develops.\n"
    ),
    "Travel/japan_trip_itinerary.html": (
        "<html><head><title>Itinerary</title></head><body><h1>Itinerary</h1><ul>"
        "<li>Day 1 (10 Oct): land at Narita, check in at Hotel Gracery, Shinjuku.</li>"
        "<li>Day 3: Shinkansen to Kyoto.</li><li>Day 4: Fushimi Inari and Arashiyama.</li>"
        "<li>Day 6: Osaka, Dotonbori street food.</li><li>Day 8 (17 Oct): fly home from Kansai.</li>"
        "</ul><script>var tracking = 1;</script></body></html>"
    ),
    "Travel/flight_confirmation_AI306.eml": {
        "From": "Air India <noreply@airindia.example>",
        "To": "rohan.mehta@example.com",
        "Subject": "Booking confirmed - PNR K7XQ2L",
        "body": (
            "Your booking is confirmed.\n\nFlight AI 306, Delhi (DEL) to Tokyo Narita (NRT)\n"
            "Departs 10 Oct 2024 01:15, arrives 12:40 local time.\nPassengers: Rohan Mehta, Meera Mehta\n"
            "Baggage: 2 x 23 kg checked, 7 kg cabin per passenger."
        ),
    },
    "Travel/visa_documents_list.md": (
        "# Japan tourist visa (VFS) - documents\n\n- Passport and one 45x45 mm photo\n"
        "- Bank statements for the last 6 months\n- ITR for the last 3 years\n"
        "- Flight and hotel bookings\n- Cover letter with the travel dates\n"
    ),
    "Home/wifi_router_notes.md": (
        "# TP-Link Archer C6\n\n- Admin page: http://192.168.0.1 (user: admin)\n"
        "- Reset: hold the reset button for 10 seconds\n- 5 GHz network name: Mehta_5G\n"
        "- ISP: ACT Fibernet, customer ID 1234567\n"
    ),
    "Home/society_maintenance_notice.pdf": (
        "Green Park Society - Notice\nMaintenance charges are revised to Rs. 3.5 per sq ft from April 2024.\n"
        "Water tank cleaning on 22 June; water supply off from 10 am to 2 pm.\n"
        "Annual general meeting (AGM) on 30 June at 6 pm in the clubhouse."
    ),
    "Home/home_inventory.xlsx": [
        ["Item", "Brand", "Purchased", "Price (Rs.)"],
        ["Refrigerator", "Samsung", "18-05-2023", "34,990"],
        ["Washing machine", "LG", "02-11-2022", "31,500"],
        ["Laptop", "Dell XPS 13", "20-01-2024", "1,12,499"],
        ["Air conditioner", "Voltas 1.5 ton", "10-04-2023", "38,000"],
    ],
}


def write(folder: Path) -> Path:
    import docx
    import openpyxl
    import pymupdf
    from pptx import Presentation

    for rel, content in DOCS.items():
        path = folder / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        ext = path.suffix.lower()
        if ext == ".pdf":
            doc = pymupdf.open()
            for page_text in str(content).split("\f"):
                doc.new_page().insert_textbox(pymupdf.Rect(72, 72, 540, 770), page_text, fontsize=11)
            doc.save(path)
            doc.close()
        elif ext == ".docx":
            d = docx.Document()
            for para in str(content).split("\n\n"):
                d.add_paragraph(para)
            d.save(path)
        elif ext == ".xlsx":
            wb = openpyxl.Workbook()
            for row in content:
                wb.active.append(row)
            wb.save(path)
        elif ext == ".pptx":
            prs = Presentation()
            for title, body in content:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = title
                slide.placeholders[1].text = body
            prs.save(path)
        elif ext == ".eml":
            msg = EmailMessage()
            for header in ("From", "To", "Subject"):
                msg[header] = content[header]
            msg.set_content(content["body"])
            path.write_bytes(bytes(msg))
        else:
            path.write_text(str(content), encoding="utf-8")
    return folder


if __name__ == "__main__":
    print("wrote", write(Path(sys.argv[1])))
