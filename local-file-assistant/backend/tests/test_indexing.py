import docx
import openpyxl
from pptx import Presentation

from app.core import indexer
from app.core.chunking import split_words
from app.core.parsers import PARSERS
from app.db import sqlite_fts


def test_split_words_overlaps():
    words = [f"w{i}" for i in range(600)]
    pieces = split_words(" ".join(words), size=250, overlap=40)
    assert len(pieces) == 3
    assert pieces[1].split()[0] == "w210"  # starts 40 words before the end of piece 1
    assert pieces[-1].split()[-1] == "w599"


def test_docx_is_cited_by_part_not_page(corpus):
    chunks = PARSERS[".docx"](corpus / "meeting_notes.docx")
    assert chunks and {c.loc_kind for c in chunks} == {"part"}
    assert [c.loc_no for c in chunks] == list(range(1, len(chunks) + 1))


def test_pdf_pages(corpus):
    chunks = PARSERS[".pdf"](corpus / "invoice_notes.pdf")
    assert [(c.loc_kind, c.loc_no) for c in chunks] == [("page", 1), ("page", 2)]


def test_pptx_and_xlsx_parsers(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    slide.placeholders[1].text = "Budget reserve 20000"
    prs.save(tmp_path / "deck.pptx")
    wb = openpyxl.Workbook()
    wb.active.title = "Q3"
    wb.active.append(["Item", "Amount"])
    wb.active.append(["Launch reserve", 20000])
    wb.save(tmp_path / "budget.xlsx")

    slides = PARSERS[".pptx"](tmp_path / "deck.pptx")
    assert slides[0].loc_kind == "slide" and "Budget reserve" in slides[0].text
    sheets = PARSERS[".xlsx"](tmp_path / "budget.xlsx")
    assert sheets[0].loc_kind == "sheet" and "row 2: Launch reserve | 20000" in sheets[0].text


def test_rescan_skips_unchanged_and_prunes_deleted(corpus, db_paths):
    db, vdb = db_paths
    first = indexer.index_folder(corpus, db, vdb)
    assert first["indexed"] == 2 and first["failed"] == 0
    second = indexer.index_folder(corpus, db, vdb)
    assert second["skipped"] == 2 and second["indexed"] == 0

    (corpus / "meeting_notes.docx").unlink()
    third = indexer.index_folder(corpus, db, vdb)
    assert third["removed"] == 1
    conn = sqlite_fts.connect(db)
    try:
        assert [p for p in sqlite_fts.indexed_paths(conn) if p.endswith("meeting_notes.docx")] == []
    finally:
        conn.close()


def test_broken_file_is_reported_with_a_reason(corpus, db_paths):
    (corpus / "broken.pdf").write_bytes(b"not a pdf")
    db, vdb = db_paths
    indexer.RECENT_ERRORS.clear()
    result = indexer.index_folder(corpus, db, vdb)
    assert result["failed"] == 1
    assert indexer.RECENT_ERRORS[0]["path"].endswith("broken.pdf") and indexer.RECENT_ERRORS[0]["error"]


def test_embedder_down_keeps_keyword_index_and_retries_later(corpus, db_paths, monkeypatch):
    def down(*_a, **_k):
        raise ConnectionError("ollama not running")

    monkeypatch.setattr(indexer, "embed", down)
    db, vdb = db_paths
    result = indexer.index_folder(corpus, db, vdb)
    assert result["indexed"] == 2 and result["semantic"] is False
    conn = sqlite_fts.connect(db)
    try:
        assert all(f["hash"] == "" for f in sqlite_fts.list_files(conn))  # retried on next scan
    finally:
        conn.close()


def test_office_lock_files_are_ignored(corpus):
    doc = docx.Document()
    doc.add_paragraph("lock")
    doc.save(corpus / "~$meeting_notes.docx")
    assert not indexer.is_indexable(corpus / "~$meeting_notes.docx")


def test_text_html_and_email_parsers(tmp_path):
    from email.message import EmailMessage

    from app.core.parsers import text_parser

    (tmp_path / "notes.md").write_text("# Router\nAdmin page: 192.168.0.1", encoding="utf-8")
    (tmp_path / "old.txt").write_bytes("Café receipt".encode("cp1252"))  # not valid UTF-8
    (tmp_path / "page.html").write_text(
        "<html><head><title>Trip</title><style>p{}</style></head><body><p>Kyoto &amp; Osaka</p>"
        "<script>var secret = 1;</script></body></html>",
        encoding="utf-8",
    )
    msg = EmailMessage()
    msg["Subject"], msg["From"] = "Booking confirmed - PNR K7XQ2L", "Air India <a@example.com>"
    msg.set_content("Baggage: 2 x 23 kg")
    (tmp_path / "flight.eml").write_bytes(bytes(msg))

    text = {p.name: "\n".join(c.text for c in PARSERS[p.suffix](p)) for p in tmp_path.iterdir()}
    assert "192.168.0.1" in text["notes.md"]
    assert text["old.txt"] == "Café receipt"
    assert "Kyoto & Osaka" in text["page.html"] and "Trip" in text["page.html"]
    assert "secret" not in text["page.html"] and "p{}" not in text["page.html"]
    assert "PNR K7XQ2L" in text["flight.eml"] and "2 x 23 kg" in text["flight.eml"]
    assert all(c.loc_kind == "part" for c in PARSERS[".eml"](tmp_path / "flight.eml"))

    # The size cap can cut a multi-byte character in half; that must not turn the file into cp1252 mojibake.
    big = tmp_path / "big.txt"
    big.write_bytes(b"a" * (text_parser.MAX_TEXT_BYTES - 1) + "é".encode("utf-8"))
    assert set(text_parser._read(big)) == {"a"}


def test_embedder_sees_file_and_folder_but_stored_text_does_not(tmp_path, db_paths, monkeypatch):
    from app.db import vector_store

    db, vdb = db_paths
    seen: list[str] = []
    fake = indexer.embed
    monkeypatch.setattr(indexer, "embed", lambda texts, **kw: seen.extend(texts) or fake(texts))
    root = tmp_path / "docs"
    (root / "Insurance").mkdir(parents=True)
    (root / "Insurance" / "car_policy.txt").write_text("Motor package policy", encoding="utf-8")

    indexer.index_folder(root, db, vdb)
    assert seen == ["File: car_policy.txt\nFolder: Insurance\n\nMotor package policy"]
    stored = vector_store.open_table(vector_store.connect(vdb)).to_arrow().column("text").to_pylist()
    assert stored == ["Motor package policy"]  # what answers and citations are checked against


def test_file_whose_vectors_failed_to_store_is_retried(corpus, db_paths, monkeypatch):
    from app.db import vector_store

    db, vdb = db_paths
    real = vector_store.upsert

    def disk_full(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(vector_store, "upsert", disk_full)
    assert indexer.index_folder(corpus, db, vdb)["failed"] == 2

    # The hash used to be stored before the vectors, so these files were skipped for good.
    monkeypatch.setattr(vector_store, "upsert", real)
    again = indexer.index_folder(corpus, db, vdb)
    assert again["indexed"] == 2 and again["skipped"] == 0
    table = vector_store.open_table(vector_store.connect(vdb))
    assert len(set(table.to_arrow().column("path").to_pylist())) == 2


def test_reindex_and_delete_touch_only_that_files_chunks(tmp_path, db_paths):
    from pathlib import Path

    db, vdb = db_paths
    root = (tmp_path / "docs").resolve()
    root.mkdir()
    (root / "a.txt").write_text("alpha one", encoding="utf-8")
    (root / "b.txt").write_text("bravo two", encoding="utf-8")
    indexer.index_folder(root, db, vdb)
    (root / "a.txt").write_text("alpha changed entirely", encoding="utf-8")
    indexer.index_folder(root, db, vdb)

    conn = sqlite_fts.connect(db)
    try:
        texts = {Path(r[0]).name: r[1] for r in conn.execute("SELECT path, text FROM chunks")}
        assert texts == {"a.txt": "alpha changed entirely", "b.txt": "bravo two"}  # replaced, not duplicated
        assert {Path(p).name: t for p, t in sqlite_fts.first_texts(conn, str(root)).items()} == texts
        assert [Path(r[0]).name for r in conn.execute("SELECT path FROM chunks WHERE chunks MATCH 'bravo'")] == ["b.txt"]
        sqlite_fts.delete_file(conn, root / "a.txt")
        assert [Path(r[0]).name for r in conn.execute("SELECT path FROM chunks")] == ["b.txt"]
    finally:
        conn.close()
